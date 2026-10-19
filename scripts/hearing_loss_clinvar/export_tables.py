#!/usr/bin/env python3
"""Export standalone analysis tables from fixed result files."""
import argparse
import csv
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from export_disease_context import generate_disease_context_display

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results"
OUTDIR = ROOT / "tables"

def hpo_route_rows():
    """Describe mutually exclusive inclusion-route unions at VariationID level."""
    route_names = {"HPO_disease_id": "HPO-listed disease identifier", "HPO_disease_name_exact": "HPO-listed disease name",
                   "HPO_hearing_term": "HPO hearing term", "expanded_text": "Expanded phenotype text"}
    expanded_path = RESULTS / "expanded/variant_aggregate.tsv"
    disease_path = RESULTS / "disease_supported/variant_aggregate.tsv"
    with expanded_path.open(encoding="utf-8", newline="") as f:
        expanded_ids = {row["variation_id"] for row in csv.DictReader(f, delimiter="\t")}
    counts = Counter()
    with disease_path.open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f, delimiter="\t"):
            routes = tuple(sorted(filter(None, row["inclusion_reasons"].split("|"))))
            if not routes:
                raise RuntimeError(f"Missing inclusion route for VariationID {row['variation_id']}")
            if any(route not in route_names for route in routes):
                raise RuntimeError(f"Unknown inclusion route for VariationID {row['variation_id']}: {routes}")
            counts[(routes, row["variation_id"] in expanded_ids)] += 1
    routes_seen = {route for route, _ in counts}
    result = []
    for route in sorted(routes_seen, key=lambda r: (-sum(counts[(r, e)] for e in [False, True]), r)):
        overlap = counts[(route, True)]
        additional = counts[(route, False)]
        result.append({"selection_routes": " + ".join(route_names[name] for name in route), "n_variation_ids": overlap + additional,
                       "n_in_expanded_text": overlap, "n_additional_to_expanded_text": additional})
    total_overlap = sum(row["n_in_expanded_text"] for row in result)
    if total_overlap != len(expanded_ids):
        raise RuntimeError("Expanded-text VariationIDs are not a subset of the disease-annotation definition")
    result.append({"selection_routes": "All selected VariationIDs", "n_variation_ids": sum(row["n_variation_ids"] for row in result),
                   "n_in_expanded_text": total_overlap,
                   "n_additional_to_expanded_text": sum(row["n_additional_to_expanded_text"] for row in result)})
    return result


def reader_header(field):
    overrides = {
        "review_tier_proxy": "Review stars",
        "n_variants": "VariationIDs",
        "n_variation_ids": "VariationIDs",
        "n_unique_variants": "Annotated VariationIDs",
        "n::pathogenic_or_likely_pathogenic": "P/LP",
        "n::vus": "VUS",
        "n::conflicting": "Conflicting",
        "n::benign_or_likely_benign": "Benign/likely benign",
        "n::other": "Other classifications",
        "n::unknown": "Unknown classifications",
        "n_plp_review_stars_ge_2": "P/LP with ≥2 stars",
        "plp_review_stars_ge_2_proportion": "P/LP ≥2-star proportion",
        "n_plp_review_stars_ge_2_gene_symbols": "Annotated genes with P/LP at ≥2 stars",
        "n_plp": "P/LP VariationIDs",
        "n_vus": "VUS VariationIDs",
        "n_conflicting": "Conflicting VariationIDs",
        "n_benign": "Benign/likely benign VariationIDs",
        "n_other": "Other VariationIDs",
        "plp_review_stars_ge_2_rate": "P/LP ≥2-star proportion",
        "n_annotated_variants": "Annotated VariationIDs",
        "n_all_annotation_variants": "All annotated VariationIDs",
        "n_single_gene_variants": "Single-gene VariationIDs",
        "n_multi_gene_variants": "Multi-gene VariationIDs",
        "n_no_gene_variants": "VariationIDs without parsed genes",
        "n_all_conflicting": "Conflicting VariationIDs, all annotations",
        "n_single_gene_conflicting": "Conflicting VariationIDs, single-gene subset",
        "n_all_uncertainty": "Uncertain/conflicting VariationIDs, all annotations",
        "n_single_gene_uncertainty": "Uncertain/conflicting VariationIDs, single-gene subset",
        "uncertainty_proportion_rank_n_ge_10": "Uncertainty-rate rank (genes with ≥10 annotations)",
        "variant_gene_memberships": "Variant–gene memberships",
        "gene_variant_memberships": "Variant–gene memberships",
        "sum_of_gene_level_counts": "Sum of gene-level VariationID counts",
        "n_unavailable_review_status": "VariationIDs with unavailable review status",
        "bucket_variant_level": "Variant classification group",
        "review_tier_proxy_max": "Highest review stars",
        "gene_parse_status": "Gene annotation status",
        "classification_group": "Classification group",
        "left": "First selection definition",
        "right": "Second selection definition",
        "n_left": "VariationIDs in first definition",
        "n_right": "VariationIDs in second definition",
        "n_intersection": "VariationIDs in both definitions",
        "n_left_only": "VariationIDs in first definition only",
        "n_right_only": "VariationIDs in second definition only",
        "gene_symbol": "Gene symbol",
        "variation_id": "ClinVar VariationID",
        "variation_ids": "VariationIDs",
        "unique_variation_ids": "VariationIDs",
    }
    if field in overrides:
        return overrides[field]
    acronyms = {"plp": "P/LP", "vus": "VUS", "scv": "SCV", "scvs": "SCVs", "vcv": "VCV",
                "hpo": "HPO", "id": "ID", "ids": "IDs", "qc": "QC", "rcv": "RCV", "rcvs": "RCVs"}
    tokens = field.replace("::", "_").split("_")
    if tokens[0] == "n":
        tokens = ["Number", "of"] + tokens[1:]
    label = " ".join(acronyms.get(token, token) for token in tokens)
    return label[0].upper() + label[1:] if label else field


def reader_value(field, text):
    if field in {"definition", "selection_definition", "left", "right"}:
        return {"primary": "Hearing-loss phrase", "expanded": "Expanded text",
                "disease_supported": "Text plus HPO", "curated_gene": "Gene comparator"}.get(text, text)
    if field in {"classification_group", "bucket_variant_level"}:
        return {"pathogenic_or_likely_pathogenic": "P/LP", "vus": "VUS", "conflicting": "Conflicting",
                "benign_or_likely_benign": "Benign/likely benign", "other": "Other", "unknown": "Unknown"}.get(text, text)
    if field == "inclusion_reasons":
        names = {"original_phrase": "Hearing-loss phrase", "expanded_text": "Expanded hearing-related text",
                 "HPO_disease_id": "HPO-listed disease identifier", "HPO_disease_name_exact": "HPO-listed disease name",
                 "HPO_hearing_term": "HPO hearing term", "curated_gene_annotation": "Qualifying gene annotation"}
        return "; ".join(names.get(item, item) for item in text.split("|"))
    return text


def add_codebook(wb, field_mappings):
    definitions = {
        "definition": "Selection definition used for this row.",
        "selection_definition": "Selection definition used for this row.",
        "n_variation_ids": "Number of distinct ClinVar VariationIDs.",
        "unique_variation_ids": "Number of distinct ClinVar VariationIDs.",
        "n_variants": "Number of distinct ClinVar VariationIDs in the indicated group.",
        "n_unique_variants": "Number of distinct VariationIDs annotated to this gene.",
        "n_gene_symbols": "Distinct explicitly named gene symbols on selected records, including coannotated symbols outside the selection reference; unresolved labels are excluded.",
        "n_multi_gene_variants": "VariationIDs annotated to more than one parsed gene symbol.",
        "n_single_gene_variants": "VariationIDs annotated to exactly one parsed gene symbol.",
        "n_no_gene_variants": "VariationIDs with no parsed gene symbol.",
        "n_plp_review_stars_ge_2": "P/LP VariationIDs with at least two official ClinVar review stars.",
        "plp_review_stars_ge_2_proportion": "P/LP at two or more stars divided by distinct VariationIDs in this definition.",
        "n_plp_review_stars_ge_2_gene_symbols": "All explicitly named gene symbols on selected P/LP records at two or more stars, including coannotated symbols outside the selection reference.",
        "n_expert_panel": "VariationIDs with expert-panel review status.",
        "n_practice_guideline": "VariationIDs with practice-guideline review status.",
        "n_vus": "VariationIDs in the uncertain-significance classification group.",
        "n_conflicting": "VariationIDs in the conflicting-classifications group.",
        "n_unavailable_review_status": "VariationIDs lacking a mapped official review-star status.",
        "n_variants_with_participating_germline_scvs": "VariationIDs listing at least one participating germline SCV.",
        "sum_variant_participating_germline_scv_counts": "Sum of listed participating germline SCV counts across VariationIDs.",
        "review_tier_proxy": "Official ClinVar review stars; zero is a valid category, unlike a missing value.",
        "n::pathogenic_or_likely_pathogenic": "VariationIDs in the pathogenic/likely pathogenic group.",
        "n::vus": "VariationIDs in the uncertain-significance group.",
        "n::conflicting": "VariationIDs in the conflicting-classifications group.",
        "n::benign_or_likely_benign": "VariationIDs in the benign/likely benign group.",
        "n::other": "VariationIDs in other classification groups.",
        "n::unknown": "VariationIDs with an unmapped classification group.",
        "gene_symbol": "Parsed ClinVar gene symbol; annotation does not establish causality.",
        "n_plp": "Annotated VariationIDs in the pathogenic/likely pathogenic group.",
        "n_vus": "Annotated VariationIDs in the uncertain-significance group.",
        "n_conflicting": "Annotated VariationIDs in the conflicting-classifications group.",
        "n_benign": "Annotated VariationIDs in the benign/likely benign group.",
        "n_other": "Annotated VariationIDs in other classification groups.",
        "uncertainty_burden": "Uncertain and conflicting annotated VariationIDs combined.",
        "conflict_rate": "Conflicting VariationIDs divided by annotated VariationIDs for this gene.",
        "vus_rate": "Uncertain-significance VariationIDs divided by annotated VariationIDs for this gene.",
        "plp_review_stars_ge_2_rate": "P/LP VariationIDs at two or more stars divided by annotated VariationIDs for this gene.",
        "n_annotated_variants": "VariationIDs with this parsed gene annotation.",
        "curation_match": "Gene–disease curation match category in the reference comparison.",
        "curation_sources": "Reference curation source labels for matched gene–disease relationships.",
        "classifications": "Reference gene–disease validity classification labels.",
        "disease_ids": "Disease identifiers in matched gene–disease curation rows.",
        "submitters": "Submitting groups in matched gene–disease curation rows.",
        "n_reference_curation_rows": "Number of matched gene–disease reference rows.",
        "n_all_annotation_variants": "VariationIDs with this gene annotation, including multi-gene annotations.",
        "n_all_conflicting": "Conflicting VariationIDs with this gene annotation.",
        "n_single_gene_conflicting": "Conflicting VariationIDs annotated to this gene only.",
        "n_all_uncertainty": "Uncertain or conflicting VariationIDs with this gene annotation.",
        "n_single_gene_uncertainty": "Uncertain or conflicting single-gene VariationIDs.",
        "uncertainty_proportion": "Uncertainty count divided by the indicated gene-annotation denominator.",
        "uncertainty_count_rank": "Gene rank by uncertain or conflicting VariationID count.",
        "uncertainty_proportion_rank_n_ge_10": "Gene rank by uncertainty proportion among genes with at least 10 annotations.",
        "summary_row_policy": "Rule used to summarize differing source classification rows.",
        "variation_ids": "Number of distinct VariationIDs in the definition.",
        "records_with_differing_source_classifications": "VariationIDs with differing source classification strings.",
        "records_with_mixed_source_categories": "VariationIDs with source classifications in multiple broad categories.",
        "records_changing_classification_category": "VariationIDs whose broad category changes under this summary rule.",
        "variation_id": "ClinVar VariationID.",
        "gene_symbols": "Parsed gene symbols for this VariationID; may contain multiple annotations.",
        "bucket_variant_level": "Broad classification group summarized at VariationID level.",
        "review_tier_proxy_max": "Highest official ClinVar review-star tier in the selected source rows.",
        "condition_names": "Condition names from selected source rows; may contain multiple conditions.",
        "condition_ids": "Condition identifiers from selected source rows; may contain multiple conditions.",
        "inclusion_reasons": "Union of matching selection routes across rows for this VariationID; routes need not refer to the same condition row.",
        "scv_accession": "Individual ClinVar submission accession.",
        "scv_version": "Version of the SCV accession retrieved for this example.",
        "submitter": "Submitting organization for the SCV example.",
        "organization_category": "Organization category recorded for the submission.",
        "classification": "Classification recorded in the cited SCV or VCV example.",
        "review_status": "ClinVar review-status text for the cited record.",
        "review_stars": "Official review stars for the cited SCV example.",
        "date_last_evaluated": "Clinical evaluation date recorded for the SCV example.",
        "date_updated": "Record update date for the SCV example.",
        "date_submitted": "Submission date for the SCV example.",
        "contributes_to_aggregate": "Whether the SCV contributes to the shown aggregate classification.",
        "vcv_accession": "ClinVar aggregate variant accession.",
        "vcv_version": "Version of the VCV accession retrieved for this example.",
        "name": "ClinVar record name.",
        "last_updated": "Last update date of the VCV record.",
        "number_of_submissions": "Submissions listed for the VCV example.",
        "number_of_submitters": "Submitters listed for the VCV example.",
        "official_review_stars": "Official ClinVar review stars for the VariationID.",
        "source_gene_annotation": "ClinVar GeneSymbol annotation as supplied by the source.",
        "parsed_gene_symbols": "Parsed gene symbols retained for gene-level summaries.",
        "gene_parse_status": "Status of GeneSymbol parsing for this VariationID.",
        "unresolved_annotation_count": "Number of GeneSymbol annotations not parsed as gene symbols.",
        "classification_group": "Broad ClinVar classification group.",
        "left": "First selection definition in the pairwise overlap.",
        "right": "Second selection definition in the pairwise overlap.",
        "n_left": "Distinct VariationIDs in the first definition.",
        "n_right": "Distinct VariationIDs in the second definition.",
        "n_intersection": "Distinct VariationIDs present in both definitions.",
        "n_left_only": "Distinct VariationIDs present only in the first definition.",
        "n_right_only": "Distinct VariationIDs present only in the second definition.",
        "variant_gene_memberships": "Total VariationID–parsed-gene memberships; multi-gene variants contribute more than one membership.",
        "sum_of_gene_level_counts": "Sum of per-gene VariationID counts; equals membership count, not distinct-variant count.",
        "selection_routes": "Mutually exclusive combination of route flags after VariationID-level union; not a same-condition attribution.",
        "n_in_expanded_text": "VariationIDs also present in the expanded-text definition.",
        "n_additional_to_expanded_text": "VariationIDs absent from the expanded-text definition.",
        "Selection routes": "Mutually exclusive route-flag combinations pooled across selected rows at VariationID level; flags need not refer to the same condition.",
        "VariationIDs": "Number of distinct selected VariationIDs in this route combination.",
        "Also in expanded-text set": "VariationIDs also present in the expanded-text definition.",
        "Additional to expanded-text set": "VariationIDs absent from the expanded-text definition.",
        "Disease identifier": "HPO-selected disease identifier matched on a ClinVar source row.",
        "Disease name": "Disease label supplied by the versioned HPO reference.",
        "Newly selected VariationIDs": "Distinct VariationIDs absent from the expanded-text definition and present on source rows matching this disease ID; rows overlap and must not be summed.",
        "Evidence threshold": "ClinGen gene–disease grade threshold for this comparator; broad includes Definitive, Strong, Moderate, and Limited, whereas Moderate or higher includes Definitive, Strong, and Moderate.",
        "Qualifying genes": "Number of distinct parsed gene symbols meeting the stated ClinGen grade threshold.",
        "Selected VariationIDs": "Distinct VariationIDs annotated to at least one qualifying gene.",
        "VariationIDs also in phrase set": "Selected VariationIDs also present in the original hearing-loss phrase definition.",
        "Expert-panel VariationIDs": "Selected VariationIDs with an expert-panel review status.",
        "P/LP VariationIDs at two or more stars": "Selected VariationIDs with P/LP classification and at least two review stars on the same source row.",
    }
    ws = wb.create_sheet("Codebook")
    ws.append(["Worksheet", "Column", "Meaning and unit"])
    for title, fields in field_mappings.items():
        for source_field, display_field in fields:
            field = display_field if title in {"S13_HPORoutes", "S14_DiseaseIDs", "S15_ClinGenGrades"} else source_field
            if field not in definitions:
                raise RuntimeError(f"Missing codebook definition: {title}.{field}")
            ws.append([title, display_field, definitions[field]])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions


def workbook():
    dest = OUTDIR; dest.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); intro = wb.active; intro.title = "Readme"
    intro.append(["Supporting tables", "Source/definition"])
    intro.append(["ClinVar input", "Accessed 2026-05-23; source and URL are listed in docs/DATA_MANIFEST.tsv; fixed-input SHA256 is recorded in run_metadata.json"])
    intro.append(["Counting", "Distinct VariationID records; a variant counts once in each annotated gene. Gene totals are not additive."])
    intro.append(["Gene fields", "Gene counts include all explicitly named symbols on selected records, including coannotations outside the gene reference. Records with descriptive labels or missing gene fields remain in record totals; S10 includes all 13 phrase-selected records without a parsed symbol."])
    intro.append(["Reference versions", "HPO v2026-09-01; ClinGen gene-validity export accessed 2026-10-13."])
    intro.append(["S13 units", "Combinations of selection criteria matched by each distinct VariationID."])
    intro.append(["S14 units", "Distinct newly selected VariationIDs per matched disease identifier. A record may appear under several identifiers, so rows are nonadditive."])
    intro.append(["S15 scope", "All condition annotations for genes qualifying at the indicated ClinGen evidence threshold."])
    tables = {
        "S1_Coverage": "coverage_summary.tsv", "S2_Stars": "primary/utility_overall_by_tier.tsv",
        "S3_Genes": "primary/utility_gene_all.tsv", "S4_Validity": "gene_disease_validity_comparison.tsv",
        "S5_Sensitivity": "gene_annotation_sensitivity.tsv", "S6_Collapse": "aggregation_policy_sensitivity.tsv",
        "S7_ExpertsAdded": "additional_disease_expert_variants.tsv", "S8_SCVExamples": "retrieved_scv_examples.tsv",
        "S9_VCVExamples": "retrieved_vcv_examples.tsv", "S10_Annotation": "primary/variant_aggregate.tsv",
        "S11_Overlap": "definition_overlaps.tsv", "S12_Accounting": "numerical_verification.tsv"}
    descriptions = {
        "S1_Coverage": "Coverage of the four selection definitions.",
        "S2_Stars": "Classification composition by official ClinVar review stars.",
        "S3_Genes": "Primary-selection gene profiles.",
        "S4_Validity": "Gene-disease validity evidence for annotated genes.",
        "S5_Sensitivity": "Gene-membership sensitivity analysis.",
        "S6_Collapse": "Classification consistency across VariationID-level summary-row policies.",
        "S7_ExpertsAdded": "Additional expert-panel records selected through disease annotations.",
        "S8_SCVExamples": "Individual submitted classifications for the targeted records.",
        "S9_VCVExamples": "Aggregate VCV metadata for the targeted records.",
        "S10_Annotation": "Records with criteria-provided conflicts, multiple annotated genes, descriptive gene labels, or no named gene symbol. Of the 13 phrase-selected records without a parsed symbol, six have descriptive labels and seven have source gene field '-'.",
        "S11_Overlap": "Pairwise overlap among selection definitions.",
        "S12_Accounting": "VariationID and gene-membership totals by selection definition.",
        "S13_HPORoutes": "Mutually exclusive selection-route combinations for the disease-annotation definition.",
        "S14_DiseaseIDs": "Leading HPO-selected disease IDs matched on source rows among newly selected VariationIDs.",
        "S15_ClinGenGrades": "Gene-based comparator sensitivity to ClinGen validity-grade thresholds.",
    }
    expected_rows = {}
    field_mappings = {}
    for title, relative in tables.items():
        path = RESULTS / relative
        with path.open() as f:
            reader = csv.reader(f, delimiter="\t"); rows = list(reader)
        if title == "S6_Collapse":
            rows = [[row[i] for i in [0, 1, 2, 3, 4, 5]] for row in rows]
            rows[0] = ["selection_definition", "summary_row_policy", "variation_ids",
                       "records_with_differing_source_classifications", "records_with_mixed_source_categories",
                       "records_changing_classification_category"]
            policy_labels = {
                "fixed_precedence": "Fixed classification precedence",
                "require_same_bucket_else_other": "Require a consistent classification category",
                "highest_review_source_rows": "Use source rows with the highest review status",
            }
            for row in rows[1:]:
                row[1] = policy_labels.get(row[1], row[1])
        elif title == "S10_Annotation":
            records = [dict(zip(rows[0], row)) for row in rows[1:]]
            rows = [["variation_id", "official_review_stars", "source_gene_annotation",
                     "parsed_gene_symbols", "gene_parse_status", "unresolved_annotation_count",
                     "classification_group", "review_status"]]
            for record in records:
                if not ((record["bucket_variant_level"] == "conflicting" and record["review_tier_proxy_max"] == "1")
                        or int(record["gene_symbol_set_size"]) > 1 or int(record["unresolved_gene_annotation_count"]) > 0
                        or int(record["gene_symbol_set_size"]) == 0 or "subset" in record["gene_parse_statuses"]):
                    continue
                status_labels = {
                    "resolved": "Named gene symbols",
                    "resolved_explicit_subset": "Named genes in a subset",
                    "resolved_explicit_subset_unpaired_ids": "Named genes in a subset; GeneID pairing unavailable",
                    "unresolved_summary_label": "Description without named genes",
                    "unresolved_non_symbol_subset": "Subset without recognizable gene symbols",
                    "unresolved_non_symbol_label": "Annotation without recognizable gene symbols",
                    "missing": "Gene annotation unavailable",
                }
                statuses = "; ".join(status_labels.get(status, status.replace("_", " ")) for status in record["gene_parse_statuses"].split("|"))
                rows.append([record["variation_id"], record["review_tier_proxy_max"], record["raw_gene_annotations"],
                             record["gene_symbols"], statuses, record["unresolved_gene_annotation_count"],
                             record["bucket_variant_level"], record["review_status_best"]])
        elif title == "S12_Accounting":
            rows = [[row[i] for i in [0, 1, 2, 3]] for row in rows]
            rows[0] = ["selection_definition", "unique_variation_ids", "variant_gene_memberships",
                       "sum_of_gene_level_counts"]
        elif title == "S2_Stars" and "4" not in {row[0] for row in rows[1:]}:
            rows.insert(1, ["4"] + ["0"] * (len(rows[0]) - 1))
        source_fields = rows[0][:]
        displayed = [reader_header(field) for field in source_fields]
        if len(set(displayed)) != len(displayed):
            raise RuntimeError(f"Duplicate displayed column names in {title}")
        rows[0] = displayed
        field_mappings[title] = list(zip(source_fields, displayed))
        ws = wb.create_sheet(title)
        for ri, row in enumerate(rows):
            values = []
            for column_index, text in enumerate(row):
                if len(text) > 32767:
                    raise RuntimeError(f"Excel cell would truncate evidence: {relative}")
                value = reader_value(source_fields[column_index], text) if ri else text
                if ri and text and title not in {"S8_SCVExamples", "S9_VCVExamples"}:
                    try: value = int(value)
                    except ValueError:
                        try: value = float(value)
                        except ValueError: pass
                values.append(value)
            ws.append(values)
        intro.append([title, descriptions[title]])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        expected_rows[title] = len(rows)
    routes = hpo_route_rows()
    ws = wb.create_sheet("S13_HPORoutes")
    ws.append(["Selection routes", "VariationIDs", "Also in expanded-text set", "Additional to expanded-text set"])
    field_mappings["S13_HPORoutes"] = list(zip(
        ["selection_routes", "n_variation_ids", "n_in_expanded_text", "n_additional_to_expanded_text"],
        [cell.value for cell in ws[1]]))
    for row in routes:
        ws.append([row[k] for k in ["selection_routes", "n_variation_ids", "n_in_expanded_text", "n_additional_to_expanded_text"]])
    intro.append(["S13_HPORoutes", descriptions["S13_HPORoutes"]])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    expected_rows["S13_HPORoutes"] = len(routes) + 1
    with (RESULTS / "hpo_disease_contributions.tsv").open(encoding="utf-8", newline="") as f:
        diseases = list(csv.DictReader(f, delimiter="\t"))
    diseases.sort(key=lambda row: (-int(row["new_variation_ids_with_matching_disease_id"]), row["disease_id"]))
    if len({row["disease_id"] for row in diseases}) != len(diseases):
        raise RuntimeError("Duplicate disease IDs in contribution table")
    ws = wb.create_sheet("S14_DiseaseIDs")
    ws.append(["Disease identifier", "Disease name", "Newly selected VariationIDs"])
    field_mappings["S14_DiseaseIDs"] = list(zip(
        ["disease_id", "disease_name", "new_variation_ids_with_matching_disease_id"],
        [cell.value for cell in ws[1]]))
    for row in diseases[:20]:
        ws.append([row["disease_id"], row["disease_name"], int(row["new_variation_ids_with_matching_disease_id"])])
    intro.append(["S14_DiseaseIDs", descriptions["S14_DiseaseIDs"]])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    expected_rows["S14_DiseaseIDs"] = min(20, len(diseases)) + 1
    with (RESULTS / "clingen_grade_sensitivity.tsv").open(encoding="utf-8", newline="") as f:
        grade_rows = list(csv.DictReader(f, delimiter="\t"))
    grade_names = {"broad": "Broad (includes Definitive, Strong, Moderate, and Limited)", "moderate_or_higher": "Moderate, Strong, or Definitive"}
    if {row["evidence_threshold"] for row in grade_rows} != set(grade_names) or len(grade_rows) != len(grade_names):
        raise RuntimeError("ClinGen grade sensitivity table must contain both comparator definitions")
    ws = wb.create_sheet("S15_ClinGenGrades")
    ws.append(["Evidence threshold", "Qualifying genes", "Selected VariationIDs", "VariationIDs also in phrase set",
               "Expert-panel VariationIDs", "P/LP VariationIDs at two or more stars"])
    field_mappings["S15_ClinGenGrades"] = list(zip(
        ["evidence_threshold", "qualifying_gene_symbols", "selected_variation_ids",
         "variation_ids_shared_with_primary_phrase", "expert_panel_records", "plp_review_stars_ge_2_records"],
        [cell.value for cell in ws[1]]))
    for row in sorted(grade_rows, key=lambda r: list(grade_names).index(r["evidence_threshold"])):
        ws.append([grade_names[row["evidence_threshold"]], int(row["qualifying_gene_symbols"]),
                   int(row["selected_variation_ids"]), int(row["variation_ids_shared_with_primary_phrase"]),
                   int(row["expert_panel_records"]), int(row["plp_review_stars_ge_2_records"])])
    intro.append(["S15_ClinGenGrades", descriptions["S15_ClinGenGrades"]])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    expected_rows["S15_ClinGenGrades"] = 3
    add_codebook(wb, field_mappings)
    for ws in wb:
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244761")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ws.columns:
            length = max(len(str(c.value or "")) for c in column[:100])
            ws.column_dimensions[get_column_letter(column[0].column)].width = min(65, max(12, length + 2))
        for row in ws.iter_rows():
            line_count = 1
            for cell in row:
                width = ws.column_dimensions[cell.column_letter].width
                capacity = max(8, int(width * 0.85))
                wrapped = sum(max(1, (len(part) + capacity - 1) // capacity)
                              for part in str(cell.value or "").split("\n"))
                line_count = max(line_count, wrapped)
            ws.row_dimensions[row[0].row].height = min(240, max(20, line_count * 14 + 4))
    with tempfile.TemporaryDirectory(prefix=".workbook-", dir=dest) as stage:
        staged = Path(stage) / "Supplementary_Tables.xlsx"
        wb.save(staged)
        check = load_workbook(staged, read_only=True, data_only=True)
        try:
            if len(check.sheetnames) != len(tables) + 5: raise RuntimeError("Workbook sheet mismatch")
            for name, n in expected_rows.items():
                if check[name].max_row != n: raise RuntimeError(f"Workbook row mismatch: {name}")
        finally:
            check.close()
        staged.replace(dest / "Supplementary_Tables.xlsx")

def table1_rows(results_dir):
    """Create the record-coverage table from the saved coverage summary."""
    path = results_dir / "coverage_summary.tsv"
    with path.open(encoding="utf-8", newline="") as stream:
        rows = list(csv.DictReader(stream, delimiter="\t"))
    labels = {
        "primary": "Primary hearing-loss phrase",
        "expanded": "Expanded text",
        "disease_supported": "Expanded text plus HPO-linked disease selection",
        "curated_gene": "Gene-annotation comparator",
    }
    expected_order = ["primary", "expanded", "disease_supported", "curated_gene"]
    by_definition = {row["definition"]: row for row in rows}
    if set(by_definition) != set(expected_order):
        raise RuntimeError(f"Unexpected definitions in {path}: {sorted(by_definition)}")
    headers = [
        "Selection definition",
        "Variation records",
        "Distinct annotated gene symbols",
        "Multi-gene records",
        "P/LP records with review status ≥2",
        "Annotated genes in that P/LP subset",
        "Expert-panel records",
    ]
    source_fields = [
        "n_variation_ids", "n_gene_symbols", "n_multi_gene_variants",
        "n_plp_review_stars_ge_2", "n_plp_review_stars_ge_2_gene_symbols",
        "n_expert_panel",
    ]
    expected = {
        "primary": (18384, 298, 62, 1810, 124, 173),
        "expanded": (25661, 467, 86, 2456, 179, 205),
        "disease_supported": (473704, 2971, 962, 29427, 1215, 2428),
        "curated_gene": (162683, 1484, 1861, 4745, 243, 241),
    }
    output = [headers]
    for definition in expected_order:
        source = by_definition[definition]
        values = tuple(int(source[field]) for field in source_fields)
        if values != expected[definition]:
            raise RuntimeError(f"Coverage values do not match the pinned table for {definition}: {values}")
        output.append([labels[definition], *values])
    return output


def table2_rows(display_rows):
    """Create the five-row disease-context table from its traceable source."""
    expected_keys = [
        "ush2a_2353", "cftr_practice_guideline_all", "hpo_disease_1",
        "hpo_disease_2", "hpo_disease_3",
    ]
    if [row["key"] for row in display_rows] != expected_keys:
        raise RuntimeError("Disease-context display rows do not match the manuscript table order")
    headers = ["Record or disease annotation", "Records", "Evidence in the fixed export"]
    evidence_by_stars = {3: "expert panel", 4: "practice guideline"}
    output = [headers]
    for row in display_rows:
        stars = row["review_stars"]
        if stars:
            stars = int(stars)
            if stars not in evidence_by_stars:
                raise RuntimeError(f"Unexpected review-star count for {row['key']}: {stars}")
            evidence = f"{row['classification_label']}; {stars} review stars ({evidence_by_stars[stars]})"
        else:
            evidence = "Matching disease identifier"
        output.append([row["row_label"], f"{int(row['n_variation_ids']):,}", evidence])
    return output


def table3_rows(results_dir):
    """Create the primary-definition review-star table from saved tier results."""
    path = results_dir / "primary/utility_overall_by_tier.tsv"
    with path.open(encoding="utf-8", newline="") as stream:
        source_rows = list(csv.DictReader(stream, delimiter="\t"))
    required = {
        "review_tier_proxy", "n_variants", "n::pathogenic_or_likely_pathogenic",
        "n::vus", "n::conflicting", "n::benign_or_likely_benign",
        "n::other", "n::unknown",
    }
    if not source_rows or not required.issubset(source_rows[0]):
        raise RuntimeError(f"Missing required fields in {path}")
    by_tier = {int(row["review_tier_proxy"]): row for row in source_rows}
    if len(by_tier) != len(source_rows) or not set(by_tier).issubset({0, 1, 2, 3, 4}):
        raise RuntimeError(f"Unexpected or duplicate review-star tier in {path}")
    if 4 not in by_tier:
        by_tier[4] = {
            "review_tier_proxy": "4",
            "n_variants": "0",
            "n::pathogenic_or_likely_pathogenic": "0",
            "n::vus": "0",
            "n::conflicting": "0",
            "n::benign_or_likely_benign": "0",
            "n::other": "0",
            "n::unknown": "0",
        }
    headers = ["Review stars", "Variation records", "P/LP", "VUS", "Conflicting", "B/LB", "Other"]
    output = [headers]
    for tier in range(5):
        row = by_tier[tier]
        values = [
            int(row["n_variants"]),
            int(row["n::pathogenic_or_likely_pathogenic"]),
            int(row["n::vus"]),
            int(row["n::conflicting"]),
            int(row["n::benign_or_likely_benign"]),
            int(row["n::other"]) + int(row["n::unknown"]),
        ]
        tier_label = {3: "3: expert panel", 4: "4: practice guideline"}.get(tier, tier)
        output.append([tier_label, *values])
    expected = [
        (0, 1018, 611, 325, 5, 9, 68),
        (1, 10727, 2393, 3896, 2800, 1636, 2),
        (2, 6466, 1722, 2644, 0, 2100, 0),
        ("3: expert panel", 173, 88, 43, 0, 41, 1),
        ("4: practice guideline", 0, 0, 0, 0, 0, 0),
    ]
    observed = [tuple(row) for row in output[1:]]
    if observed != expected:
        raise RuntimeError(f"Primary review-tier values do not match the pinned table: {observed}")
    return output


def write_tsv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=path.parent,
                                     prefix=f".{path.name}.", suffix=".partial", delete=False) as stream:
        staged = Path(stream.name)
        writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
        writer.writerows(rows)
    staged.replace(path)


def main_tables_workbook(path, tables):
    """Write an optional three-sheet workbook containing the main table exports."""
    wb = Workbook()
    for index, (sheet_name, rows) in enumerate(tables):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = sheet_name
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244761")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[get_column_letter(column[0].column)].width = min(65, max(12, width + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".main-tables-", dir=path.parent) as stage:
        staged = Path(stage) / path.name
        wb.save(staged)
        staged.replace(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", type=Path, default=ROOT / "results",
                        help="directory containing the saved analysis result tables")
    parser.add_argument("--outdir", type=Path, default=ROOT / "tables",
                        help="directory for exported workbooks and TSV files")
    parser.add_argument("--main-workbook", action="store_true",
                        help="also write Main_Tables.xlsx with Table1, Table2, and Table3 sheets")
    args = parser.parse_args()
    global RESULTS, OUTDIR
    RESULTS = args.results.expanduser().resolve()
    OUTDIR = args.outdir.expanduser().resolve()
    OUTDIR.mkdir(parents=True, exist_ok=True)
    workbook()
    display_rows, _, _ = generate_disease_context_display(RESULTS, RESULTS)
    table1 = table1_rows(RESULTS)
    table2 = table2_rows(display_rows)
    table3 = table3_rows(RESULTS)
    write_tsv(OUTDIR / "Table1.tsv", table1)
    write_tsv(OUTDIR / "Table2.tsv", table2)
    write_tsv(OUTDIR / "Table3.tsv", table3)
    if args.main_workbook:
        main_tables_workbook(OUTDIR / "Main_Tables.xlsx", [
            ("Table1", table1), ("Table2", table2), ("Table3", table3),
        ])


if __name__ == "__main__":
    main()
